"""Lectura/validacion del Excel de casos de prueba y armado del XML de steps."""

from __future__ import annotations

import re
from html import escape
from pathlib import Path

import pandas as pd

from qa_toolkit.core.errors import ExcelError

COLUMNAS_REQUERIDAS = [
    "ID caso de prueba",
    "Módulo",
    "Funcionalidad",
    "Nombre del caso de prueba",
    "Descripción del caso de prueba",
    "Instrucciones de ejecución",
    "Resultado esperado",
]


def limpiar_valor(valor) -> str:
    """Celda -> texto limpio ('' si esta vacia)."""
    if pd.isna(valor):
        return ""
    return str(valor).strip()


def quitar_numeracion(texto: str) -> str:
    """'1. Abrir' / '2) Abrir' -> 'Abrir'."""
    return re.sub(r"^\s*\d+\s*[\.\)\-:]\s*", "", texto).strip()


def construir_steps_xml(instrucciones: str, resultados_esperados: str) -> str:
    """Un ActionStep por linea de instrucciones, emparejado por posicion con el resultado."""
    acciones = [quitar_numeracion(linea) for linea in instrucciones.splitlines() if linea.strip()]
    resultados = [quitar_numeracion(linea) for linea in resultados_esperados.splitlines() if linea.strip()]
    if not acciones:
        acciones = ["Ejecutar el caso de prueba según la descripción."]
    while len(resultados) < len(acciones):
        resultados.append("")
    partes = [f'<steps id="0" last="{len(acciones) + 1}">']
    for step_id, (accion, resultado) in enumerate(zip(acciones, resultados, strict=False), start=2):
        partes.append(
            f'<step id="{step_id}" type="ActionStep">'
            f'<parameterizedString isformatted="true">{escape(accion)}</parameterizedString>'
            f'<parameterizedString isformatted="true">{escape(resultado)}</parameterizedString>'
            "<description/></step>"
        )
    partes.append("</steps>")
    return "".join(partes)


def cargar_excel(excel_file: str, sheet_name: str, solo_primer_caso: bool = False) -> pd.DataFrame:
    """Lee y valida el Excel. Lanza ExcelError ante cualquier problema."""
    if not excel_file:
        raise ExcelError("No se indicó el archivo Excel.")
    if not Path(excel_file).is_file():
        raise ExcelError(f"No se encontró el Excel:\n{excel_file}")
    try:
        df = pd.read_excel(excel_file, sheet_name=sheet_name, engine="openpyxl", dtype=str)
    except ValueError as e:
        raise ExcelError(f"Error leyendo la hoja '{sheet_name}': {e}") from e
    except Exception as e:  # noqa: BLE001
        raise ExcelError(f"Error leyendo el Excel: {e}") from e

    faltantes = [c for c in COLUMNAS_REQUERIDAS if c not in df.columns]
    if faltantes:
        raise ExcelError("Faltan columnas requeridas:\n" + "\n".join(f"- {c}" for c in faltantes))

    df = df[COLUMNAS_REQUERIDAS].copy()
    df = df[df["ID caso de prueba"].notna()].copy()
    df["ID caso de prueba"] = df["ID caso de prueba"].astype(str).str.strip()
    df = df[df["ID caso de prueba"] != ""].copy()

    duplicados = df[df["ID caso de prueba"].duplicated(keep=False)]
    if not duplicados.empty:
        ids = sorted(duplicados["ID caso de prueba"].unique())
        raise ExcelError("Hay IDs duplicados en el Excel:\n" + "\n".join(f"- {i}" for i in ids))
    if df.empty:
        raise ExcelError("El Excel no tiene casos de prueba (todas las filas sin ID).")
    return df.head(1) if solo_primer_caso else df


def generar_plantilla(destino: str) -> str:
    """Crea un Excel de plantilla con instrucciones y una fila de ejemplo."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Casos de prueba"
    ws.append(COLUMNAS_REQUERIDAS)
    ws.append(
        [
            "CP-001",
            "Facturación",
            "Alta de factura",
            "Crear factura con datos válidos",
            "Verifica que se pueda emitir una factura con todos los datos obligatorios.",
            "1. Ingresar al módulo Facturación\n2. Completar los datos obligatorios\n3. Presionar Emitir",
            "Se muestra el módulo\nLos datos se aceptan sin errores\nLa factura se emite y aparece en el listado",
        ]
    )
    relleno = PatternFill("solid", fgColor="1F3864")
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = relleno
        celda.alignment = Alignment(vertical="center", wrap_text=True)
    for col, ancho in zip("ABCDEFG", (16, 18, 22, 34, 44, 52, 52), strict=True):
        ws.column_dimensions[col].width = ancho
    for celda in ws[2]:
        celda.alignment = Alignment(vertical="top", wrap_text=True)

    ayuda = wb.create_sheet("Instrucciones")
    for linea in (
        "Una fila = un Test Case. No repetir el 'ID caso de prueba'.",
        "'Instrucciones de ejecución': un paso por línea (Alt+Enter). La numeración inicial es opcional.",
        "'Resultado esperado': una línea por paso, en el mismo orden que las instrucciones.",
        "Se pueden agregar columnas propias a la derecha: se ignoran.",
    ):
        ayuda.append([linea])
    ayuda.column_dimensions["A"].width = 110
    wb.save(destino)
    return destino
